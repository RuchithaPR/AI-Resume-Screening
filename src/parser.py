"""Text extraction from PDF, DOCX and TXT files."""
from __future__ import annotations

import io
import re
from pathlib import Path

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


def _clean(text: str) -> str:
    text = text.replace("\x00", " ").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t\u00a0]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _read_pdf(data: bytes) -> str:
    import pdfplumber

    pages = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            pages.append(page.extract_text() or "")
    return "\n".join(pages)


def _read_docx(data: bytes) -> str:
    import docx  # python-docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def extract_text_from_bytes(data: bytes, filename: str) -> str:
    """Extract plain text from the bytes of an uploaded file."""
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        raw = _read_pdf(data)
    elif ext == ".docx":
        raw = _read_docx(data)
    elif ext in (".txt", ".md"):
        raw = data.decode("utf-8", errors="ignore")
    else:
        raise ValueError(f"Unsupported file type '{ext}' ({filename}). Use PDF, DOCX or TXT.")
    return _clean(raw)


def extract_text(path: str | Path) -> str:
    """Extract plain text from a file on disk."""
    path = Path(path)
    return extract_text_from_bytes(path.read_bytes(), path.name)
